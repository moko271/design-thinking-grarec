# -*- coding: utf-8 -*-
"""
app.py：デザイン思考ワークシート（/design）を足す（2026-10-13）
  ・/design ルート（kadai.html を開く。画面側で ?sheet=design と同じ扱いになる）
  ・PowerPoint書き出し（sheet: 'design'）：
      見出しの表を「テーマ｜チームメンバ」にし、左上の「課題1-1」の所に各スライドのフェーズ名（①さぐる など）
      各スライドにエリアの枠と見出し・説明を描いてから（いちばん後ろ）、カードとテキストボックスを置く
何度実行しても同じ結果になる。
"""
import sys

p = sys.argv[1] if len(sys.argv) > 1 else "app.py"
s = open(p, encoding="utf-8").read()
if "_kadai_add_area(" in s:
    print("already patched"); sys.exit(0)


def rep(old, new):
    global s
    assert s.count(old) == 1, old[:80]
    s = s.replace(old, new)


rep('''@app.route("/kadai")
def kadai():''', '''@app.route("/design")
def design_sheet():
    # デザイン思考ワークシート（①さぐる ②きづく ③ひらめく）。画面は課題シートと同じ
    return send_from_directory("static", "kadai.html")


@app.route("/kadai")
def kadai():''')

rep('''    slides_in = [x for x in slides_in if isinstance(x, dict)][:30] or [{"cards": [], "texts": []}]
''', '''    slides_in = [x for x in slides_in if isinstance(x, dict)][:30] or [{"cards": [], "texts": []}]
    is_design = data.get("sheet") == "design"
    theme = str(data.get("theme", "")).strip()[:80]
''')

rep('''    # 2枚目以降：見出しを書き込んだ1枚目の図形をそのままコピーする（レイアウトの空のタイトル枠は消す）''',
    '''    # デザイン思考ワークシート：見出しを「テーマ｜チームメンバ」に、左上はフェーズ名（あとで各スライドに入れる）
    if is_design:
        for shape in slide.shapes:
            if shape.has_text_frame and not shape.is_placeholder and "課題" in shape.text_frame.text:
                _kadai_set_text(shape.text_frame, "__PHASE__")
            if shape.has_table:
                tb = shape.table
                total = sum(col.width for col in tb.columns)
                tb.columns[0].width = int(total * 0.38)
                tb.columns[1].width = total - int(total * 0.38)
                _kadai_set_text(tb.cell(0, 0).text_frame, "テーマ")
                _kadai_set_text(tb.cell(1, 0).text_frame, theme)
                _kadai_set_text(tb.cell(0, 1).text_frame, "チームメンバ")

    # 2枚目以降：見出しを書き込んだ1枚目の図形をそのままコピーする（レイアウトの空のタイトル枠は消す）''')

rep('''    for pg, sd in zip(pages, slides_in):
        for c in (sd.get("cards") or [])[:80]:''', '''    for pg, sd in zip(pages, slides_in):
        if is_design:
            for shp in pg.shapes:
                if shp.has_text_frame and not shp.is_placeholder and "__PHASE__" in shp.text_frame.text:
                    _kadai_set_text(shp.text_frame, str(sd.get("label") or ""))
                    for r in shp.text_frame.paragraphs[0].runs:
                        r.font.size = Pt(18)
            for ar in (sd.get("areas") or [])[:12]:
                _kadai_add_area(pg, ar, Emu, Pt, RGBColor, MSO_SHAPE)
        for c in (sd.get("cards") or [])[:80]:''')

rep('''    name = f"課題{_safe_filename_part(kadai_no)}" + (f"_{_safe_filename_part(group)}班" if group else "") + ".pptx"''',
    '''    name = (f"デザイン思考" if is_design else f"課題{_safe_filename_part(kadai_no)}") + (f"_{_safe_filename_part(group)}班" if group else "") + ".pptx"''')

helpers = '''
def _kadai_set_text(tf, text):
    """書式（最初の文字の大きさ・色など）を残したまま、文字だけ入れ替える"""
    paras = tf.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    runs = paras[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        paras[0].add_run().text = text


def _kadai_color(v, default=(0xDD, 0xDD, 0xDD)):
    """'#9c27b0' または 'rgba(243,229,245,.75)'（白の上に重ねた色にする）を (r, g, b) に"""
    v = str(v or "").strip()
    try:
        if v.startswith("#") and len(v) == 7:
            return tuple(int(v[i:i + 2], 16) for i in (1, 3, 5))
        if v.startswith("rgba(") or v.startswith("rgb("):
            parts = [float(x) for x in v[v.index("(") + 1:v.index(")")].split(",")]
            a = parts[3] if len(parts) > 3 else 1.0
            return tuple(int(round(c * a + 255 * (1 - a))) for c in parts[:3])
    except (ValueError, IndexError):
        pass
    return default


def _kadai_add_area(slide, a, Emu, Pt, RGBColor, MSO_SHAPE):
    """デザイン思考ワークシートのエリア（枠・見出し・説明）"""
    b = KADAI_BODY
    x, y = _kadai_pct(a.get("x", 0)), _kadai_pct(a.get("y", 0))
    w, h = _kadai_pct(a.get("w", 30), 1.0), _kadai_pct(a.get("h", 30), 1.0)
    if None in (x, y, w, h):
        return
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(b["left"] + b["width"] * x / 100)),
                                 Emu(int(b["top"] + b["height"] * y / 100)),
                                 Emu(int(b["width"] * w / 100)), Emu(int(b["height"] * h / 100)))
    shp.adjustments[0] = 0.03
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(*_kadai_color(a.get("b"), (0xF5, 0xF5, 0xF5)))
    shp.line.color.rgb = RGBColor(*_kadai_color(a.get("c"), (0x99, 0x99, 0x99)))
    shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    from pptx.enum.text import MSO_ANCHOR
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(72000))
    from pptx.enum.text import PP_ALIGN
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    r = tf.paragraphs[0].add_run()
    r.text = str(a.get("label", ""))[:40]
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    hint = str(a.get("hint", ""))[:120]
    if hint:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = hint
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
'''
anchor = '''def _kadai_add_card(slide, c, Emu, Pt, RGBColor, MSO_SHAPE, MSO_ANCHOR):'''
rep(anchor, helpers.lstrip("\n") + "\n\n" + anchor)
open(p, "w", encoding="utf-8").write(s)
print("patched")
