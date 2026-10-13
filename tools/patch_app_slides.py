# -*- coding: utf-8 -*-
"""
app.py：/api/export_kadai_pptx を複数スライド＋テキストボックスに対応させる（2026-10-13）
  受け取る形（新）: slides: [{cards: [{keyword, x, y}], texts: [{text, x, y, w, size}]}, ...]
  これまでの形（cards だけ）も受け付ける（1枚目に置く）
  2枚目以降は、見出し（課題番号・リーダ学番・メンバー）を書き込んだ1枚目の図形をコピーして作る
何度実行しても同じ結果になる。
"""
import sys

p = sys.argv[1] if len(sys.argv) > 1 else "app.py"
s = open(p, encoding="utf-8").read()
if "_kadai_add_card(" in s:
    print("already patched"); sys.exit(0)


def rep(old, new):
    global s
    assert s.count(old) == 1, old[:80]
    s = s.replace(old, new)


rep('''    members = data.get("members") or []
    cards = data.get("cards") or []
    group = str(data.get("group_number", "")).strip()[:10]
''', '''    members = data.get("members") or []
    cards = data.get("cards") or []
    group = str(data.get("group_number", "")).strip()[:10]
    slides_in = data.get("slides")
    if not isinstance(slides_in, list) or not slides_in:
        slides_in = [{"cards": cards, "texts": []}]
    slides_in = [x for x in slides_in if isinstance(x, dict)][:30] or [{"cards": [], "texts": []}]
''')

old_start = '''    # カード（画面上の位置 x,y は自由スペースに対する%）
    b = KADAI_BODY
    for c in cards[:80]:'''
old_end = '''        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    buf = io.BytesIO()'''
i0 = s.index(old_start)
i1 = s.index(old_end) + len(old_end)
new_block = '''    # 2枚目以降：見出しを書き込んだ1枚目の図形をそのままコピーする（レイアウトの空のタイトル枠は消す）
    import copy as _copy
    pages = [slide]
    for _ in range(len(slides_in) - 1):
        ns = prs.slides.add_slide(slide.slide_layout)
        for shp in list(ns.shapes):
            shp._element.getparent().remove(shp._element)
        for shp in slide.shapes:
            ns.shapes._spTree.append(_copy.deepcopy(shp._element))
        pages.append(ns)
    # ページ番号（フィールドなのでPowerPointで開くと自動で直るが、念のため表示も合わせる）
    for no, pg in enumerate(pages, 1):
        for shp in pg.shapes:
            if shp.is_placeholder and shp.has_text_frame:
                for t in shp._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"):
                    t.text = str(no)

    # カードとテキストボックス（画面上の位置 x,y,w は自由スペースに対する%）
    for pg, sd in zip(pages, slides_in):
        for c in (sd.get("cards") or [])[:80]:
            _kadai_add_card(pg, c, Emu, Pt, RGBColor, MSO_SHAPE, MSO_ANCHOR)
        for t in (sd.get("texts") or [])[:80]:
            _kadai_add_text(pg, t, Emu, Pt, RGBColor)

    buf = io.BytesIO()'''
s = s[:i0] + new_block + s[i1:]

helpers = '''

def _kadai_pct(v, lo=0.0, hi=100.0):
    try:
        return max(lo, min(float(v), hi))
    except (TypeError, ValueError):
        return None


def _kadai_add_card(slide, c, Emu, Pt, RGBColor, MSO_SHAPE, MSO_ANCHOR):
    b = KADAI_BODY
    x, y = _kadai_pct(c.get("x", 0)), _kadai_pct(c.get("y", 0))
    if x is None or y is None:
        return
    left = min(int(b["left"] + b["width"] * x / 100), b["left"] + b["width"] - KADAI_CARD_W)
    top = min(int(b["top"] + b["height"] * y / 100), b["top"] + b["height"] - KADAI_CARD_H)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(top), Emu(KADAI_CARD_W), Emu(KADAI_CARD_H))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xFF, 0xF9, 0xD9)
    box.line.color.rgb = RGBColor(0xE6, 0xD9, 0x8C)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(54000))
    run = tf.paragraphs[0].add_run()
    run.text = str(c.get("keyword", ""))[:60]
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)


def _kadai_add_text(slide, t, Emu, Pt, RGBColor):
    """ふつうのPowerPointのテキストボックスと同じ（枠線・塗りなし、幅固定で折り返し、高さは文字に合わせる）"""
    from pptx.enum.text import MSO_AUTO_SIZE
    b = KADAI_BODY
    text = str(t.get("text", "")).replace("\\r", "")[:2000]
    if not text.strip():
        return
    x, y = _kadai_pct(t.get("x", 0)), _kadai_pct(t.get("y", 0))
    w = _kadai_pct(t.get("w", 30), 3.0, 100.0)
    if x is None or y is None or w is None:
        return
    try:
        size = max(8, min(int(float(t.get("size", 18))), 72))
    except (TypeError, ValueError):
        size = 18
    lines = text.split("\\n")
    left = int(b["left"] + b["width"] * x / 100)
    top = int(b["top"] + b["height"] * y / 100)
    width = int(b["width"] * w / 100)
    height = int(Pt(size).emu * 1.3 * len(lines) + 91440)
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
'''
anchor = '''@app.route("/api/export_kadai_pptx", methods=["POST"])'''
rep(anchor, helpers.lstrip("\n") + "\n\n" + anchor)
open(p, "w", encoding="utf-8").write(s)
print("patched")
