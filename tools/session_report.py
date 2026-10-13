# -*- coding: utf-8 -*-
"""
授業1回分のセッション（/kadai・/design）を振り返るレポートを作る。APIは使わない。
grarec_light フォルダで:
    python tools/session_report.py <セッションのJSON> [--summary <班のまとめ>]
      <班のまとめ> … 班が最後にまとめたもの。.pptx（書き出したもの・提出したもの）か .txt（1行に1項目）
→ 同じ場所に <セッション名>_report.md と <セッション名>_cards.csv

レポートの中身
  1. 概要：録音の長さ、文字起こしの区間（失敗・無音）、カードが届いた回と枚数
  2. カードの行方：届いた／置いた／書き直した／消した／置かれず下がった、手で書いたカード、テキストボックス
  3. 話のバランス：お知らせの判定と、実際に出た時刻・閉じた時刻
  4. 班のまとめとAIカードの比較（--summary があるとき）
       班のまとめの各項目に、いちばん近いAIカード（文字の重なり＝バイグラムJaccard）を並べる
       0.25以上を「近い」、0.12〜0.25を「少し近い」とする。粗い判定なので、最後は目で確かめること
"""
import argparse
import csv
import json
import os
import re
import unicodedata
from collections import Counter
from datetime import datetime


def norm(s):
    return re.sub(r"[\s、。！？!?「」『』・（）()\[\]]", "", unicodedata.normalize("NFKC", s or "")).lower()


def bigrams(s):
    s = norm(s)
    return {s[i:i + 2] for i in range(len(s) - 1)} if len(s) > 1 else ({s} if s else set())


def sim(a, b):
    A, B = bigrams(a), bigrams(b)
    return len(A & B) / len(A | B) if A and B else 0.0


def ts(x):
    if isinstance(x, (int, float)):
        return x / 1000
    try:
        return datetime.fromisoformat(str(x).replace("Z", "+00:00")).timestamp()
    except ValueError:
        return None


def summary_items(path):
    if path.lower().endswith(".pptx"):
        from pptx import Presentation
        items = []
        for no, sl in enumerate(Presentation(path).slides, 1):
            for shp in sl.shapes:
                if shp.is_placeholder or shp.has_table or not shp.has_text_frame:
                    continue
                t = shp.text_frame.text.strip()
                if not t or re.fullmatch(r"課題\s*[\d\-]+|[①②③④⑤].{0,6}|__PHASE__", t):
                    continue
                items.append((no, t))
        return items
    return [(None, line.strip()) for line in open(path, encoding="utf-8") if line.strip()]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("session")
    ap.add_argument("--summary")
    a = ap.parse_args()
    d = json.load(open(a.session, encoding="utf-8"))
    ev = d.get("events", [])
    t0 = (d.get("recStartedAt") or 0) / 1000 or (ts(ev[0]["t"]) if ev else 0)
    mm = lambda t: f"{(t - t0) / 60:5.1f}分" if t is not None else "  –  "
    cards = d.get("cards", [])
    texts = d.get("texts", [])
    slides = d.get("slides", [{"id": "s1"}])
    sno = {s["id"]: i + 1 for i, s in enumerate(slides)}
    cnt = Counter(e["ev"] for e in ev)
    out = []
    w = out.append
    name = os.path.splitext(os.path.basename(a.session))[0]
    w(f"# 振り返りレポート：{name}\n")
    w(f"- 班：{d.get('group_number', '')}　シート：{d.get('sheet') or d.get('template', 'kadai')}　テーマ／課題：{d.get('task_hint') or d.get('kadai_no', '')}")
    rs = [ts(e["t"]) for e in ev if e["ev"] == "rec_start"]
    re_ = [ts(e["t"]) for e in ev if e["ev"] == "rec_stop"]
    if rs:
        w(f"- 録音：{datetime.fromtimestamp(rs[0]).strftime('%H:%M')} から" + (f" {((re_[-1] - rs[0]) / 60):.0f}分" if re_ else "（止めた記録なし）"))
    tr = d.get("transcript", [])
    w(f"- 文字起こしの区間：{len(tr)}（話者分離あり {sum(1 for x in tr if x.get('diarized'))}）、"
      f"失敗 {cnt['seg_fail']}、無音で送らず {cnt['seg_empty']}、話者分離できず {cnt['seg_nodiarize']}、"
      f"文字数 {sum(len(x.get('text', '')) for x in tr)}")
    if cnt["save_fail"]:
        w(f"- 保存の失敗：{cnt['save_fail']} 回（回復 {cnt['save_recovered']} 回）")
    w("")

    # カードが届いた回
    w("## 1. カードが届いた回\n")
    w("| 時刻 | きっかけ | 枚数 | 文字数 |\n|---|---|---|---|")
    for e in ev:
        if e["ev"] == "cards":
            w(f"| {mm(ts(e['t']))} | {'ボタン' if e.get('trigger') == 'now' else '自動'} | {e.get('count', '')} | {e.get('chars', '')} |")
    w("")

    # カードの行方
    ai = [c for c in cards if c.get("source") == "ai"]
    manual = [c for c in cards if c.get("source") == "manual"]
    placed_ids = {e.get("id") for e in ev if e["ev"] == "place"}
    edited_ids = {e.get("id") for e in ev if e["ev"] == "edit"}
    deleted_ids = {e.get("id") for e in ev if e["ev"] == "delete"}
    w("## 2. カードの行方\n")
    w(f"- AIカード：届いた {len(ai)} 枚 → シートに置いた {sum(1 for c in ai if c['id'] in placed_ids)}、"
      f"今もシートにある {sum(1 for c in ai if c.get('place') == 'sheet')}、書き直した {sum(1 for c in ai if c['id'] in edited_ids)}、"
      f"消した {sum(1 for c in ai if c['id'] in deleted_ids)}、置かれず下がった {sum(1 for c in ai if c.get('place') == 'archived')}")
    w(f"- 種類（届いたAIカード）：" + "、".join(f"{k} {v}" for k, v in Counter(c.get('kind') for c in ai).most_common()))
    w(f"- 手で書いたカード：{len(manual)} 枚　テキストボックス：{len(texts)} 個　スライド：{len(slides)} 枚")
    w(f"- 「これまでのカード」：開いた {cnt['history_open']} 回、戻した {cnt['history_restore']} 回、そこから置いた {sum(1 for e in ev if e['ev'] == 'place' and e.get('via') == 'history')} 回")
    w("")
    w("| 届いた | 種類 | カード | 引用 | 行方 |\n|---|---|---|---|---|")
    for c in sorted(cards, key=lambda c: c.get("addedAt", "")):
        fate = {"sheet": f"シート{sno.get(c.get('slide'), 1)}枚目" + (f"（{c['area']}）" if c.get("area") else ""),
                "archived": "下がった", "tray": "トレイ", "deleted": "消した"}.get(c.get("place"), c.get("place"))
        if c["id"] in edited_ids:
            fate += "・書き直し"
        w(f"| {mm(ts(c.get('addedAt')))} | {'手書き' if c.get('source') == 'manual' else c.get('kind', '')} | {c.get('keyword', '')} | {c.get('quote', '')[:30]} | {fate} |")
    if texts:
        w("\n**班が書いたテキストボックス**\n")
        for t in texts:
            w(f"- （スライド{sno.get(t.get('slide'), 1)}）" + t.get("text", "").replace("\n", " / "))
    w("")

    # 話のバランス
    w("## 3. 話のバランスとお知らせ\n")
    bj = [e for e in ev if e["ev"] in ("balance_judge", "balance_notice", "balance_hidden", "balance_dismiss")]
    if not bj:
        w("- 記録なし（声の登録をしていない、または判定前に終了）")
    for e in bj:
        extra = ""
        if e["ev"] == "balance_judge":
            extra = f"{e.get('from')} → {e.get('to')}　均衡度 {e.get('b')}　割合 {e.get('shares')}"
        elif e["ev"] == "balance_notice":
            extra = f"お知らせが出た（判定 {e.get('judged_at', '')[11:19]}）　割合 {e.get('shares')}"
        elif e["ev"] == "balance_hidden":
            extra = f"消えた（{e.get('reason')}）"
        else:
            extra = "「閉じる」が押された"
        w(f"- {mm(ts(e['t']))} {extra}")
    w("")

    # 班のまとめとの比較
    if a.summary:
        items = summary_items(a.summary)
        w("## 4. 班のまとめとAIカードの比較\n")
        w(f"班のまとめ：{os.path.basename(a.summary)}（{len(items)} 項目）。AIカード＝届いたAIカード全部（置かなかったものも含む）\n")
        w("| 班のまとめ | いちばん近いAIカード | 近さ | 判定 |\n|---|---|---|---|")
        near = 0
        used = set()
        for no, it in items:
            best = max(ai, key=lambda c: max(sim(it, c.get("keyword", "")), sim(it, c.get("quote", ""))), default=None)
            sc = max(sim(it, best.get("keyword", "")), sim(it, best.get("quote", ""))) if best else 0
            lab = "近い" if sc >= 0.25 else "少し近い" if sc >= 0.12 else "対応なし"
            if sc >= 0.12:
                near += 1; used.add(best["id"])
            w(f"| {('(' + str(no) + ') ') if no else ''}{it.replace(chr(10), ' / ')[:50]} | {best.get('keyword', '') if best and sc >= 0.12 else ''} | {sc:.2f} | {lab} |")
        w(f"\n- 班のまとめ {len(items)} 項目のうち、AIカードと近いものがあった：{near} 項目")
        w(f"- AIカード {len(ai)} 枚のうち、班のまとめのどれかに近かった：{len(used)} 枚")
        miss = [c for c in ai if c["id"] not in used]
        if miss:
            w("- 班のまとめに出てこなかったAIカード：" + "、".join(c.get("keyword", "") for c in miss[:30]))
    base = os.path.splitext(a.session)[0]
    open(base + "_report.md", "w", encoding="utf-8").write("\n".join(out) + "\n")
    with open(base + "_cards.csv", "w", newline="", encoding="utf-8-sig") as f:
        cw = csv.writer(f)
        cw.writerow(["届いた分", "出どころ", "種類", "カード", "引用", "置き場所", "スライド", "エリア", "書き直し", "回数"])
        for c in cards:
            cw.writerow([round(((ts(c.get("addedAt")) or t0) - t0) / 60, 1) if c.get("addedAt") else "", c.get("source"), c.get("kind"),
                         c.get("keyword"), c.get("quote"), c.get("place"), sno.get(c.get("slide"), "") if c.get("place") == "sheet" else "",
                         c.get("area") or "", "はい" if c["id"] in edited_ids else "", c.get("mention", 1)])
    print("保存:", base + "_report.md", "/", base + "_cards.csv")


if __name__ == "__main__":
    main()
