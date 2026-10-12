# -*- coding: utf-8 -*-
"""app.py に課題シート（/kadai）用の変更を入れるパッチ。何度実行しても同じ結果になる。"""
import sys

p = sys.argv[1] if len(sys.argv) > 1 else "app.py"
s = open(p, encoding="utf-8").read()


def rep(old, new, marker):
    global s
    if marker in s:
        return
    assert s.count(old) == 1, f"見つからない/複数: {old[:60]!r}"
    s = s.replace(old, new)


# ---------- 1) Whisper ヒントを画面ごとに切り替え ----------
rep('''# =========================
# フェーズごとの設定''', '''# 課題シート（/kadai）など、テーマを固定しない画面用の軽いヒント。
# 専門用語を並べるとWhisperがそれを幻覚として出力することがあるため、最小限にする。
GENERIC_WHISPER_HINT = "大学の授業でのグループワークの話し合いです。"


def whisper_hint_for(mode: str) -> str:
    return GENERIC_WHISPER_HINT if mode == "generic" else WHISPER_HINT


# =========================
# フェーズごとの設定''', "GENERIC_WHISPER_HINT =")

if "def split_and_transcribe(file_path, suffix, hint=None)" not in s:
    rep('''def split_and_transcribe(file_path, suffix):
    """5分超の音声を5分チャンクに分割してWhisperに送り、結果を結合する。"""''',
        '''def split_and_transcribe(file_path, suffix, hint=None):
    """5分超の音声を5分チャンクに分割してWhisperに送り、結果を結合する。"""
    hint = WHISPER_HINT if hint is None else hint''', "@@never@@")
    a = s.index("def split_and_transcribe")
    b = s.index('@app.route("/api/transcribe", methods=["POST"])')
    seg = s[a:b]
    s = s[:a] + seg.replace("prompt=WHISPER_HINT", "prompt=hint") + s[b:]

rep('''    raw_sid = request.form.get("session_id", "")
    safe_sid = re.sub(r'[^\\w\\-]', '_', raw_sid, flags=re.UNICODE)[:40] if raw_sid else ""

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp_path = tmp.name
        audio_file.save(tmp_path)

    try:
        text = (split_and_transcribe(tmp_path, suffix) or "").strip()
''', '''    raw_sid = request.form.get("session_id", "")
    safe_sid = re.sub(r'[^\\w\\-]', '_', raw_sid, flags=re.UNICODE)[:40] if raw_sid else ""
    hint = whisper_hint_for(request.form.get("hint", ""))

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp_path = tmp.name
        audio_file.save(tmp_path)

    try:
        text = (split_and_transcribe(tmp_path, suffix, hint=hint) or "").strip()
''', 'hint = whisper_hint_for(request.form.get("hint", ""))')

# ---------- 2) /kadai ルート ----------
rep('''@app.route("/mendan")
def mendan():
    return send_from_directory("static", "index_ai.html")
''', '''@app.route("/mendan")
def mendan():
    return send_from_directory("static", "index_ai.html")


@app.route("/kadai")
def kadai():
    # 授業の課題フォーマット（A4横・自由記述）に合わせた課題シート画面
    return send_from_directory("static", "kadai.html")
''', 'def kadai():')

# ---------- 3) カード抽出 API と PowerPoint 書き出し API ----------
KADAI_BLOCK = r'''

# =========================
# 課題シート（/kadai）: カード抽出
# =========================
# 研究目的（対話構造のリアルタイム可視化で「議論構造」「他者の意見」「自己の考え」への
# 気づきを促す）に合わせ、「重要な発言を全部」ではなく「話し合いを一歩引いて見るための
# 少数のカード」を、役割の違う種類から選ぶ。AIは新しい答えやアイデアを足さない。
CARD_KINDS = {
    "topic":      "論点",      # いま班が中心に話していること → 議論構造
    "opinion":    "意見",      # 具体的な意見・アイデア・経験 → 他者の意見／自己の考え
    "difference": "ちがい",    # 意見が分かれた点・別の見方 → 他者の意見／議論構造
    "question":   "問い",      # 出たがまだ答えが出ていない問い → 議論構造（次の一歩）
    "agreement":  "まとまり",  # 班で合意した・決まったこと → 議論構造
}
KIND_IMPORTANCE = {"question": "high", "difference": "high", "topic": "medium",
                   "agreement": "medium", "opinion": "medium"}


def build_card_prompt(memo, prev_text, existing, first_cycle, max_cards, task_hint=""):
    system_prompt = (
        "あなたは大学の授業のグループワークを「見える化」するアシスタントです。\n"
        "参加者が自分たちの話し合いを一歩引いて見て、話の流れや他の人の意見、"
        "自分の考えに気づけるように、話し合いの記録から少数のカードを選びます。\n"
        "あなたは議事録係でも、アイデアや答えを出す人でもありません。"
        "実際に話された内容だけを、話し手の言葉をなるべく残して短くします。"
    )
    existing_txt = "\n".join(f"・{k}" for k in existing[:60]) or "（まだありません）"
    if first_cycle:
        count_rule = (
            f"・今回は最初のカードです。「この班がいま何について話しているか」がわかる\n"
            f"  論点(topic)を中心に3〜{min(4, max_cards)}枚にする\n"
        )
    else:
        count_rule = (
            f"・最大{max_cards}枚。大事なものが少なければ少なくてよい（無理に埋めない）\n"
            "・同じ種類(kind)は最大2枚まで\n"
            "・問い(question)とちがい(difference)が話し合いに出ていれば、優先して1枚ずつ入れる\n"
            "  （話し合いの次の一歩や、他の人の見方に気づくきっかけになるため）\n"
        )
    user_prompt = (
        (f"【この授業の課題】\n{task_hint}\n\n" if task_hint else "")
        + "【直前の区間の発話（文脈として参照するだけ。ここからは選ばない）】\n"
        f"{prev_text or '（なし）'}\n\n"
        "【今回の区間の発話（ここからカードを選ぶ）】\n"
        f"{memo}\n\n"
        "【すでに画面に出ているカード】\n"
        f"{existing_txt}\n\n"
        "【カードの種類 kind】\n"
        "- topic（論点）：いまこの班が中心に話していること\n"
        "- opinion（意見）：誰かが出した具体的な意見・アイデア・経験\n"
        "- difference（ちがい）：意見が分かれている点、別の見方が出た点\n"
        "- question（問い）：出たけれど、まだ答えが出ていない問い・保留になっていること\n"
        "- agreement（まとまり）：班で合意した・決まったこと\n\n"
        "【選び方】\n"
        + count_rule +
        "・1回しか出ず誰も拾わなかったが、話し合いを広げられそうな視点があれば1枚入れ、lone を true にする\n"
        "・すでに出ているカードと同じ内容は出さない。ただし同じ話題がこの区間でも中心的に\n"
        "  繰り返されたときは、既存カードの文言を一字一句そのまま keyword に入れて返す\n"
        "・新しいアイデア・答え・助言・まとめの言い換えを付け足さない（話されていないことは書かない）\n\n"
        "【選ばない発言】\n"
        "・雑談、相槌、笑い、段取りや役割分担だけの発言\n"
        "・GraRec・カード・録音・画面など、道具の操作や感想についての発言\n\n"
        "【書き方】\n"
        "・keyword：10〜20文字。話し手の言葉を残す。「〜について話す」「Aさんが〜を提案」のような説明にしない\n"
        "  良い例：昼休みに食堂が混みすぎる／なぜみんな同じ時間に行くのか？\n"
        "・quote：元の発言をそのまま（40文字以内）\n"
        "・type：fact|opinion|question|idea|reaction|instruction（発話の形式）\n"
        "・thought：fact|problem|interpret|wish（発話の意図）\n"
        "・reason：このカードを選んだ理由（20文字以内）\n\n"
        "次のJSONオブジェクトだけを返す：\n"
        '{"cards":[{"keyword":"","quote":"","kind":"topic|opinion|difference|question|agreement",'
        '"lone":false,"type":"","thought":"","reason":""}]}'
    )
    return system_prompt, user_prompt


def _select_cards(raw_cards, max_cards):
    """AIの出力を検証し、枚数・種類の上限を守らせる。"""
    out, per_kind, seen = [], defaultdict(int), set()
    for c in raw_cards:
        if not isinstance(c, dict):
            continue
        kw = str(c.get("keyword", "")).strip()
        kind = str(c.get("kind", "")).strip().lower()
        if not kw or kind not in CARD_KINDS or kw in seen:
            continue
        if per_kind[kind] >= 2:
            continue
        per_kind[kind] += 1
        seen.add(kw)
        out.append({
            "keyword": kw[:40],
            "quote": str(c.get("quote", "")).strip()[:80],
            "kind": kind,
            "kind_label": CARD_KINDS[kind],
            "lone": bool(c.get("lone", False)),
            "type": str(c.get("type", "")).strip().lower(),
            "thought": str(c.get("thought", "")).strip().lower(),
            "reason": str(c.get("reason", "")).strip()[:40],
            "importance": KIND_IMPORTANCE.get(kind, "medium"),
        })
        if len(out) >= max_cards:
            break
    return out


@app.route("/api/extract_cards", methods=["POST"])
def extract_cards():
    """
    課題シート用のカード抽出。
    memo: 今回の区間の文字起こし / prev_text: 直前の区間（文脈）
    existing: すでに画面に出ているカードの文言 / first_cycle: 最初のカードか
    """
    data = request.get_json() or {}
    memo = (data.get("memo") or "").strip()
    if len(memo) < 10:
        return jsonify({"ok": True, "cards": []})
    prev_text = (data.get("prev_text") or "").strip()[-500:]
    existing = [str(k) for k in (data.get("existing") or []) if k][:60]
    first_cycle = bool(data.get("first_cycle", False))
    max_cards = max(1, min(int(data.get("max_cards", 6) or 6), 8))
    task_hint = (data.get("task_hint") or "").strip()[:300]

    system_prompt, user_prompt = build_card_prompt(memo, prev_text, existing, first_cycle, max_cards, task_hint)
    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.3,
            response_format={"type": "json_object"},
        )
    except Exception as e:
        return jsonify({"ok": False, "error": f"OpenAI API error: {e}"}), 500

    raw = response.choices[0].message.content or "{}"
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        print("extract_cards JSON parse error:", raw[:200])
        return jsonify({"ok": False, "error": "AIの応答を読み取れませんでした"}), 500
    raw_cards = parsed.get("cards", []) if isinstance(parsed, dict) else parsed
    cards = _select_cards(raw_cards if isinstance(raw_cards, list) else [], max_cards)
    return jsonify({"ok": True, "cards": cards, "first_cycle": first_cycle})


# =========================
# 課題シート（/kadai）: PowerPoint 書き出し
# =========================
KADAI_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "templates", "kadai_format.pptx")
# 元フォーマットの自由スペース（区切り線の下〜ページ番号の上）。単位はEMU
KADAI_BODY = {"left": 287557, "top": 1320000, "width": 9360000, "height": 4980000}
KADAI_CARD_W = 1783080   # シート幅の18%（画面のカード幅と同じ比率）
KADAI_CARD_H = 560000


def _safe_filename_part(t: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]', '_', t)[:40]


@app.route("/api/export_kadai_pptx", methods=["POST"])
def export_kadai_pptx():
    try:
        from pptx import Presentation
        from pptx.util import Emu, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR
    except ImportError:
        return jsonify({"ok": False, "error": "python-pptx が入っていません。pip install python-pptx を実行してください"}), 500

    data = request.get_json() or {}
    kadai_no = str(data.get("kadai_no", "")).strip()[:10] or "1-1"
    leader_id = str(data.get("leader_id", "")).strip()[:20]
    members = data.get("members") or []
    cards = data.get("cards") or []
    group = str(data.get("group_number", "")).strip()[:10]

    prs = Presentation(KADAI_TEMPLATE)
    slide = prs.slides[0]

    for shape in slide.shapes:
        # 「課題 1-1」の番号部分
        if shape.has_text_frame and "課題" in shape.text_frame.text:
            runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
            if len(runs) >= 2:
                runs[-1].text = kadai_no
        # リーダ学番・メンバ名の表
        if shape.has_table:
            table = shape.table
            c_id = table.cell(1, 0).text_frame
            c_id.paragraphs[0].text = ""
            r = c_id.paragraphs[0].add_run()
            r.text = leader_id
            tf = table.cell(1, 1).text_frame
            for extra in tf.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            para = tf.paragraphs[0]
            names = [m for m in members if isinstance(m, dict) and str(m.get("name", "")).strip()]
            for i, m in enumerate(names):
                if i:
                    sep = para.add_run()
                    sep.text = "、"
                run = para.add_run()
                run.text = str(m["name"]).strip()[:20]
                role = m.get("role", "")
                if role == "leader":
                    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
                elif role == "observer":
                    run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)

    # カード（画面上の位置 x,y は自由スペースに対する%）
    b = KADAI_BODY
    for c in cards[:80]:
        try:
            x = max(0.0, min(float(c.get("x", 0)), 100.0))
            y = max(0.0, min(float(c.get("y", 0)), 100.0))
        except (TypeError, ValueError):
            continue
        left = int(b["left"] + b["width"] * x / 100)
        top = int(b["top"] + b["height"] * y / 100)
        left = min(left, b["left"] + b["width"] - KADAI_CARD_W)
        top = min(top, b["top"] + b["height"] - KADAI_CARD_H)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(top), Emu(KADAI_CARD_W), Emu(KADAI_CARD_H))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xFF, 0xF9, 0xD9)
        box.line.color.rgb = RGBColor(0xE6, 0xD9, 0x8C)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Emu(54000))
        run = tf.paragraphs[0].add_run()
        run.text = str(c.get("keyword", ""))[:60]
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    name = f"課題{_safe_filename_part(kadai_no)}" + (f"_{_safe_filename_part(group)}班" if group else "") + ".pptx"
    cd = f"attachment; filename=\"kadai.pptx\"; filename*=UTF-8''{urlquote(name, safe='')}"
    return Response(buf.read(),
                    mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    headers={"Content-Disposition": cd})
'''

if "def extract_cards():" not in s:
    anchor = '''# =========================
# エントリポイント
# ========================='''
    assert s.count(anchor) == 1
    s = s.replace(anchor, KADAI_BLOCK.lstrip("\n") + "\n\n" + anchor)

open(p, "w", encoding="utf-8").write(s)
print("patched", p)
